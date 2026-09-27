"""Выгрузка прогноза в CSV и XLSX — требование функциональности из ТЗ."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from app.api.deps import (
    ApiError,
    adjustment_params,
    check_hours,
    check_period,
    parse_date_param,
    parse_routes_param,
)
from app.pipeline.adjust import Adjustment
from app.pipeline.aggregate import iter_rows
from app.pipeline.ingest import get_dataset
from app.pipeline.network_events import get_repository as network_repository

router = APIRouter(prefix="/export", tags=["выгрузка"])

HEADER = ["route", "date", "hour", "prediction"]


def _collect(routes, date_from, date_to, hour_from, hour_to, adjustment,
             network=None, rounded=True):
    dataset = get_dataset()
    rows = []
    for row, value in iter_rows(
        dataset.forecast_index,
        routes=routes,
        start=date_from,
        end=date_to,
        hour_from=hour_from,
        hour_to=hour_to,
        adjustment=adjustment,
        network=network,
    ):
        rows.append((row[0], row[5], row[2], round(value) if rounded else value))
    rows.sort(key=lambda item: (item[0], item[1], item[2]))
    return rows


def _filename(extension: str) -> str:
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    return f"forecast_{stamp}.{extension}"


@router.get("", summary="Выгрузка прогноза в CSV или XLSX")
def export(
    format: str = Query("csv", description="csv или xlsx"),
    routes: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    hour_from: int = Query(0, ge=0, le=23),
    hour_to: int = Query(23, ge=0, le=23),
    adjustment: Adjustment = Depends(adjustment_params),
):
    fmt = format.lower().strip()
    if fmt not in ("csv", "xlsx"):
        raise ApiError(400, "INVALID_PARAMETER", "format принимает значения csv или xlsx")
    check_hours(hour_from, hour_to)
    start = parse_date_param(date_from, "date_from")
    end = parse_date_param(date_to, "date_to")
    route_list = parse_routes_param(routes)
    # Период целиком вне прогноза — ошибка; частичный выход в CSV не
    # сообщить, файл просто строится по пересечению.
    check_period(get_dataset(), "forecast", start, end)

    rows = _collect(
        route_list,
        start,
        end,
        hour_from,
        hour_to,
        adjustment,
        # Выгрузка совпадает с тем, что на экране: с событиями сети.
        network_repository().effect(),
    )

    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
        writer.writerow(HEADER)
        writer.writerows(rows)
        content = buffer.getvalue().encode("utf-8-sig")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{_filename("csv")}"'
            },
        )

    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ApiError(
            503, "EXPORT_UNAVAILABLE", "Выгрузка в XLSX недоступна: не установлен openpyxl"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Прогноз"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(list(row))
    for column, width in zip("ABCD", (10, 14, 8, 14)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{_filename("xlsx")}"'
        },
    )


@router.get("/submission", summary="Выгрузка в формате сдачи на лидерборд")
def export_submission():
    """Полная сетка 10 маршрутов × 61 день × 24 часа без поправок.

    И без событий сети: это model_prediction в формате лидерборда.
    Без округления до целых: значения должны совпадать с текущим скоренным CSV.
    """
    rows = _collect(None, None, None, 0, 23, None, rounded=False)
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(HEADER)
    writer.writerows(rows)
    content = buffer.getvalue().encode("utf-8")
    return StreamingResponse(
        io.BytesIO(content),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="submission.csv"'},
    )
