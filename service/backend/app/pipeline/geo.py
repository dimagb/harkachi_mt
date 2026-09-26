"""Геопривязка.

Второе звено цепочки. Читает справочники маршрутов и остановок из xlsx
и строит:
  * упорядоченные списки остановок по маршруту и направлению;
  * геометрию маршрута (ломаная) для отрисовки на карте;
  * веса остановок для разложения маршрутного прогноза по точкам.

Структура справочника организаторов, лист «Порядок_с_координатами»:

    route_id route_short_name reg_num route_type trip_id trip_short_name
    direction_id start_date end_date stop_sequence stop_id actual_date
    stop_mode is_addpoint stop_name stop_lat stop_lon

Важно: связь с данными идёт по `route_short_name` — номеру маршрута, как
его видит пассажир. НЕ по `route_id`: там внутренние идентификаторы вида
4450, которых в валидациях нет.

Ограничение данных: справочник покрывает маршруты 1, 2, 3, 4, 5, 6, 7, 10,
11, 12, а в данных присутствуют 1, 7, 11, 12, 17, 25, 26, 28, 50.
Прогнозировать надо ещё маршрут 5 (запущен 16.12.2025), он в справочнике
есть. Пересечение — пять маршрутов: 1, 5, 7, 11, 12, около половины
пассажиропотока; остальные показываются без карты. Маршруты справочника
вне задачи отбрасываются при загрузке.

Разложение прогноза по остановкам ОЦЕНОЧНОЕ: целевая величина определена на
уровне «маршрут × час», привязки валидаций к остановкам в данных нет
(`place_id` — это площадка трамвайного управления, что подтверждает пример
валидаций из справочника). Веса остановок строятся эвристически:
пересадочные узлы весомее рядовых остановок.
"""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass, field

from app import config

log = logging.getLogger(__name__)

# Предпочтительные имена колонок: сначала точное совпадение, потом по слову.
EXACT = {
    "route": ("route_short_name",),
    "direction": ("direction_id",),
    "sequence": ("stop_sequence",),
    "stop_id": ("stop_id",),
    "stop_name": ("stop_name",),
    "lat": ("stop_lat", "lat", "latitude"),
    "lon": ("stop_lon", "lon", "lng", "longitude"),
}
FUZZY = {
    "lat": ("широт",),
    "lon": ("долгот",),
    "stop_name": ("остановк", "название"),
}

# Ключевые слова, по которым остановка считается пересадочным узлом.
HUB_KEYS = ("метро", "мцк", "мцд", "вокзал", "платформа")
HUB_WEIGHT = 2.2        # во сколько раз узел весомее рядовой остановки
TERMINUS_WEIGHT = 1.4   # на конечных начинается поездка, там тоже нагрузка

LAT_RANGE = (54.0, 57.5)
LON_RANGE = (35.0, 39.5)


@dataclass
class Stop:
    stop_id: str
    name: str
    route: int
    direction: int
    sequence: int
    lat: float
    lon: float
    is_hub: bool = False

    def as_dict(self) -> dict:
        return {
            "stop_id": self.stop_id,
            "name": self.name,
            "route": self.route,
            "direction": self.direction,
            "sequence": self.sequence,
            "lat": self.lat,
            "lon": self.lon,
            "is_hub": self.is_hub,
        }


@dataclass
class GeoIndex:
    # (маршрут, направление) -> упорядоченный список остановок
    by_route_direction: dict = field(default_factory=dict)
    source_files: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def routes(self) -> list:
        return sorted({route for route, _ in self.by_route_direction})

    @property
    def stops_by_route(self) -> dict:
        """Совместимость: маршрут -> остановки первого направления."""
        return {route: self.stops(route) for route in self.routes}

    def directions(self, route: int) -> list:
        return sorted(d for r, d in self.by_route_direction if r == route)

    def stops(self, route: int, direction: int | None = None) -> list:
        if direction is not None:
            return self.by_route_direction.get((route, direction), [])
        for candidate in self.directions(route):
            return self.by_route_direction.get((route, candidate), [])
        return []

    def has_geometry(self, route: int) -> bool:
        return bool(self.stops(route))

    def route_geojson(self, route: int) -> list:
        """По одной линии на каждое направление."""
        features = []
        for direction in self.directions(route):
            stops = self.by_route_direction[(route, direction)]
            if len(stops) < 2:
                continue
            features.append(
                {
                    "type": "Feature",
                    "properties": {
                        "route": route,
                        "direction": direction,
                        "stops": len(stops),
                        "from": stops[0].name,
                        "to": stops[-1].name,
                    },
                    "geometry": {
                        "type": "LineString",
                        "coordinates": [[s.lon, s.lat] for s in stops],
                    },
                }
            )
        return features

    def stop_shares(self, route: int, direction: int | None = None) -> dict:
        """Доли остановок в маршрутном объёме.

        ОЦЕНКА, а не измерение. Логика весов:
          * пересадочный узел (метро, вокзал, МЦК, МЦД) весомее рядовой
            остановки — там объективно больше посадок;
          * конечные нагружены сильнее середины;
          * остальное распределяется мягким профилем без резких перепадов.

        Профиль опирается на названия остановок из официального справочника,
        а не на произвольную формулу, поэтому его можно объяснить.
        """
        stops = self.stops(route, direction)
        if not stops:
            return {}
        if len(stops) == 1:
            return {stops[0].stop_id: 1.0}

        n = len(stops)
        weights = []
        for i, stop in enumerate(stops):
            weight = 1.0
            if stop.is_hub:
                weight *= HUB_WEIGHT
            if i == 0 or i == n - 1:
                weight *= TERMINUS_WEIGHT
            x = (i - (n - 1) / 2) / max((n - 1) / 2, 1)
            weight *= 0.85 + 0.15 * math.exp(-1.5 * x * x)
            weights.append(weight)

        total = sum(weights)
        return {s.stop_id: w / total for s, w in zip(stops, weights)}

    def hub_stops(self, route: int) -> list:
        return [s for s in self.stops(route) if s.is_hub]


# --------------------------------------------------------------- разбор xlsx


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _match_columns(header: list) -> dict:
    cells = [_norm(c) for c in header]
    found: dict = {}

    for name, variants in EXACT.items():
        for variant in variants:
            if variant in cells:
                found[name] = cells.index(variant)
                break

    for name, keys in FUZZY.items():
        if name in found:
            continue
        for idx, cell in enumerate(cells):
            if cell and any(key in cell for key in keys):
                found[name] = idx
                break

    # запасной вариант для номера маршрута: колонка со словом route,
    # но не route_id и не route_type
    if "route" not in found:
        for idx, cell in enumerate(cells):
            if "route" in cell and cell not in ("route_id", "route_type"):
                found["route"] = idx
                break
    return found


def _to_float(value) -> float | None:
    if value is None:
        return None
    try:
        return float(str(value).replace(",", ".").strip())
    except ValueError:
        return None


def _to_int(value, default=None):
    if value is None:
        return default
    digits = ""
    for ch in str(value).strip():
        if ch.isdigit():
            digits += ch
        elif digits:
            break
    try:
        return int(digits) if digits else default
    except ValueError:
        return default


def _is_hub(name: str) -> bool:
    lowered = name.lower()
    return any(key in lowered for key in HUB_KEYS)


def load() -> GeoIndex:
    index = GeoIndex()
    try:
        from openpyxl import load_workbook
    except ImportError:
        index.warnings.append("openpyxl не установлен, справочники не прочитаны")
        return index

    files = sorted(config.DATA_DIR.rglob("*.xlsx"))
    if not files:
        index.warnings.append("xlsx-справочники не найдены")
        return index

    collected: dict = {}

    for path in files:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — файл может быть любым
            index.warnings.append(f"{path.name}: не открылся ({exc})")
            continue

        used_sheet = None
        for sheet in workbook.worksheets:
            columns = None
            for row in sheet.iter_rows(values_only=True):
                if columns is None:
                    candidate = _match_columns(list(row))
                    # нужен лист, где есть и координаты, и порядок остановок
                    if {"lat", "lon", "route", "sequence"} <= set(candidate):
                        columns = candidate
                    continue

                lat = _to_float(row[columns["lat"]])
                lon = _to_float(row[columns["lon"]])
                if lat is None or lon is None:
                    continue
                if not LAT_RANGE[0] < lat < LAT_RANGE[1]:
                    continue
                if not LON_RANGE[0] < lon < LON_RANGE[1]:
                    continue

                route = _to_int(row[columns["route"]])
                if route is None:
                    continue
                direction = _to_int(
                    row[columns["direction"]] if "direction" in columns else 0,
                    default=0,
                )
                sequence = _to_int(row[columns["sequence"]], default=0)

                raw_stop_id = (
                    row[columns["stop_id"]] if "stop_id" in columns else None
                )
                stop_key = (
                    str(raw_stop_id).strip()
                    if raw_stop_id is not None
                    else f"{sequence}"
                )
                raw_name = (
                    row[columns["stop_name"]] if "stop_name" in columns else None
                )
                name = (
                    str(raw_name).strip()
                    if raw_name is not None
                    else f"Остановка {sequence}"
                )

                collected.setdefault((route, direction), []).append(
                    Stop(
                        stop_id=f"{route}-{direction}-{stop_key}",
                        name=name,
                        route=route,
                        direction=direction,
                        sequence=sequence,
                        lat=lat,
                        lon=lon,
                        is_hub=_is_hub(name),
                    )
                )

            if columns is not None:
                used_sheet = sheet.title
                break  # лист с геометрией найден, остальные не нужны

        if used_sheet:
            index.source_files.append(f"{path.name} → лист «{used_sheet}»")

    # Справочник шире задачи: в нём есть 2, 3, 4, 6, 10, которых нет среди
    # прогнозируемых маршрутов. Без фильтра /api/geometry рисовал бы на карте
    # пять чужих маршрутов — на реальном справочнике так и было.
    foreign = sorted({r for r, _ in collected} - set(config.ROUTES))
    if foreign:
        log.info(
            "Геопривязка: маршруты справочника вне задачи пропущены: %s",
            ", ".join(str(r) for r in foreign),
        )
    collected = {
        key: stops for key, stops in collected.items() if key[0] in config.ROUTES
    }

    for stops in collected.values():
        stops.sort(key=lambda s: s.sequence)
    index.by_route_direction = collected

    if not collected:
        index.warnings.append(
            "координаты не распознаны — проверьте структуру справочников"
        )
    else:
        missing = sorted(set(config.ROUTES) - set(index.routes))
        if missing:
            index.warnings.append(
                "нет геометрии для маршрутов: "
                + ", ".join(str(r) for r in missing)
            )

    log.info(
        "Геопривязка: маршрутов с координатами %d (%s), источники: %s",
        len(index.routes),
        ", ".join(str(r) for r in index.routes),
        index.source_files,
    )
    return index


_geo: GeoIndex | None = None


def get_geo() -> GeoIndex:
    global _geo
    if _geo is None:
        _geo = load()
    return _geo
